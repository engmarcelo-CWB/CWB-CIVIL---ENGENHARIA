"""Gera a planilha .xlsx do laudo com fórmulas vivas (openpyxl)."""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from textos_padrao import GRAU_RISCO, NATUREZA, ORIGEM, SISTEMAS, TIPOS

FONT = "Arial"
F_TIT = Font(name=FONT, size=13, bold=True)
F_SUB = Font(name=FONT, size=9, italic=True, color="555555")
F_HEAD = Font(name=FONT, size=9, bold=True, color="FFFFFF")
F_TXT = Font(name=FONT, size=9)
F_IN = Font(name=FONT, size=9, color="0000FF")  # entrada
F_LINK = Font(name=FONT, size=9, color="008000")  # vínculo com outra aba
F_BOLD = Font(name=FONT, size=9, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_TOT = PatternFill("solid", fgColor="D9E1F2")
FILL_PEND = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="A6A6A6")
BORDA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

MOEDA = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
QTD = '#,##0.00;-#,##0.00;"-"'
PCT = '0.00%;-0.00%;"-"'


def _titulo(ws, titulo, dados, ncols):
    ws["A1"] = titulo
    ws["A1"].font = F_TIT
    imovel = (dados.get("imovel") or {}).get("endereco", "")
    ws["A2"] = f"Laudo nº {dados.get('numero_laudo','')} — {imovel}"
    ws["A2"].font = F_SUB
    base = dados.get("base_custos") or {}
    if base:
        regime = "desonerado" if base.get("regime") == "desonerado" else "não desonerado"
        ws["A3"] = (
            f"Base de custos: {base.get('fonte','')} {base.get('uf','')} — ref. {base.get('mes_referencia','')} — {regime}"
        )
        ws["A3"].font = F_SUB


def _cab(ws, row, cabecalhos, larguras):
    for i, (h, w) in enumerate(zip(cabecalhos, larguras), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.alignment = CENTER
        c.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _cel(ws, r, c, v, font=F_TXT, fmt=None, fill=None, align=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    cell.border = BORDA
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    cell.alignment = align or WRAP
    return cell


def _pct_inf(dados):
    v = (dados.get("programa") or {}).get("percentual_informado")
    if v in (None, ""):
        return None
    try:
        v = float(str(v).replace("%", "").replace(",", "."))
    except ValueError:
        return None
    return v / 100 if v > 1 else v


def gerar_planilha(dados, calc, caminho):
    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumo"
    tipo = dados.get("tipo", "")

    # ---------------- Dados ----------------
    ws_d = wb.create_sheet("Dados")
    _titulo(ws_d, "DADOS DO LAUDO E PARÂMETROS", dados, 3)
    ws_d.column_dimensions["A"].width = 4
    ws_d.column_dimensions["B"].width = 34
    ws_d.column_dimensions["C"].width = 70
    im = dados.get("imovel") or {}
    base = dados.get("base_custos") or {}
    sol = dados.get("solicitante") or {}
    linhas = [
        ("Nº do laudo", dados.get("numero_laudo", "")),
        ("Tipo", TIPOS.get(tipo, {}).get("titulo", tipo)),
        ("Solicitante", sol.get("nome", "")),
        ("Proprietário", (dados.get("proprietario") or {}).get("nome", "")),
        ("Endereço", im.get("endereco", "")),
        ("Município/UF", im.get("municipio_uf", "")),
        ("Matrícula", im.get("matricula", "")),
        ("Tipologia / padrão", " / ".join(x for x in [im.get("tipologia", ""), im.get("padrao", "")] if x)),
        ("Sistema construtivo", im.get("sistema_construtivo") or SISTEMAS.get(im.get("sistema_construtivo_tipo"), ("", []))[0]),
        ("Programa / instrumento", " — ".join(x for x in [(dados.get("programa") or {}).get("orgao", ""), (dados.get("programa") or {}).get("instrumento", ""), (dados.get("programa") or {}).get("id_sistema", "")] if x)),
        ("Grau de risco global", GRAU_RISCO.get(dados.get("grau_risco_global") or "", "")),
        ("Área construída (m²)", im.get("area_construida", "")),
        ("Data(s) da vistoria", ", ".join((dados.get("vistoria") or {}).get("datas", []))),
        ("Base de custos", f"{base.get('fonte','')} {base.get('uf','')}".strip()),
        ("Mês de referência", base.get("mes_referencia", "")),
        ("Regime", "Desonerado" if base.get("regime") == "desonerado" else "Não desonerado"),
    ]
    r = 5
    for k, v in linhas:
        _cel(ws_d, r, 2, k, F_BOLD)
        _cel(ws_d, r, 3, v)
        r += 1
    _cel(ws_d, r, 2, "BDI", F_BOLD)
    c = _cel(ws_d, r, 3, calc["bdi"], F_IN, PCT)
    c.comment = Comment("Entrada: BDI como fração (0,25 = 25%). Todas as fórmulas de preço com BDI usam esta célula.", "Laudo")
    ref_bdi = f"Dados!$C${r}"

    # ---------------- Orçamento ----------------
    linhas_s = calc["linhas"]
    tot_row = None
    if linhas_s:
        ws = wb.create_sheet("Orçamento")
        _titulo(ws, "QUANTITATIVOS E CUSTOS — EXECUTADO, A REFAZER E A EXECUTAR", dados, 20)
        cab = [
            "Item", "Etapa", "Fonte", "Código", "Descrição", "Un",
            "Qtd. prevista", "Qtd. executada aceita", "Qtd. a refazer", "Qtd. não executada",
            "PU s/ BDI", "PU c/ BDI", "PU demolição s/ BDI",
            "Valor previsto", "Valor executado aceito", "Custo de refazer", "Custo a executar",
            "Custo para concluir", "Valor executado bruto", "Observações",
        ]
        larg = [7, 18, 9, 10, 48, 6, 11, 11, 11, 11, 12, 12, 12, 14, 14, 14, 14, 15, 14, 40]
        HR = 5
        _cab(ws, HR, cab, larg)
        r0 = HR + 1
        for i, l in enumerate(linhas_s):
            r = r0 + i
            _cel(ws, r, 1, l["item"], align=CENTER)
            _cel(ws, r, 2, l["etapa"])
            _cel(ws, r, 3, l["fonte"], align=CENTER)
            _cel(ws, r, 4, l["codigo"], align=CENTER)
            _cel(ws, r, 5, l["descricao"])
            _cel(ws, r, 6, l["unidade"], align=CENTER)
            _cel(ws, r, 7, l["prev"], F_IN, QTD)
            _cel(ws, r, 8, l["aceita"], F_IN, QTD)
            _cel(ws, r, 9, l["refazer"], F_IN, QTD)
            _cel(ws, r, 10, f"=MAX(G{r}-H{r}-I{r},0)", fmt=QTD)
            if l["pu"] is None:
                pc = _cel(ws, r, 11, None, F_IN, MOEDA, FILL_PEND)
                pc.comment = Comment("Preço pendente: preencher com o PU da tabela referencial.", "Laudo")
            else:
                _cel(ws, r, 11, l["pu"], F_IN, MOEDA)
            _cel(ws, r, 12, f'=IF(K{r}="","",K{r}*(1+{ref_bdi}))', fmt=MOEDA)
            dem = l["demolicao"]
            if dem:
                if l["pu_dem"] is None or dem.get("preco_unitario") in (None, ""):
                    _cel(ws, r, 13, None, F_IN, MOEDA, FILL_PEND)
                else:
                    _cel(ws, r, 13, l["pu_dem"], F_IN, MOEDA)
            else:
                _cel(ws, r, 13, None, F_IN, MOEDA)
            _cel(ws, r, 14, f'=IF(K{r}="","",G{r}*L{r})', fmt=MOEDA)
            _cel(ws, r, 15, f'=IF(K{r}="","",H{r}*L{r})', fmt=MOEDA)
            _cel(ws, r, 16, f'=IF(K{r}="","",I{r}*(K{r}+N(M{r}))*(1+{ref_bdi}))', fmt=MOEDA)
            _cel(ws, r, 17, f'=IF(K{r}="","",J{r}*L{r})', fmt=MOEDA)
            _cel(ws, r, 18, f'=IF(K{r}="","",P{r}+Q{r})', fmt=MOEDA)
            _cel(ws, r, 19, f'=IF(K{r}="","",(H{r}+I{r})*L{r})', fmt=MOEDA)
            obs = l["observacao"]
            if dem:
                dtxt = f"Demolição/remoção: {dem.get('fonte','')} {dem.get('codigo','')} – {dem.get('descricao','')} ({dem.get('unidade','')})".strip()
                obs = f"{dtxt}. {obs}".strip()
            _cel(ws, r, 20, obs)
        rN = r0 + len(linhas_s) - 1
        tot_row = rN + 1
        _cel(ws, tot_row, 5, "TOTAL", F_BOLD, fill=FILL_TOT)
        for col in range(1, 21):
            ws.cell(row=tot_row, column=col).fill = FILL_TOT
            ws.cell(row=tot_row, column=col).border = BORDA
        for col in (14, 15, 16, 17, 18, 19):
            L = get_column_letter(col)
            _cel(ws, tot_row, col, f"=SUM({L}{r0}:{L}{rN})", F_BOLD, MOEDA, FILL_TOT)
        ws.auto_filter.ref = f"A{HR}:T{rN}"
        orc_ini, orc_fim = r0, rN  # fixa os limites: r0/rN são reutilizados nas outras abas

        def rng(col, a=orc_ini, b=orc_fim):
            return f"'Orçamento'!${col}${a}:${col}${b}"

    # ---------------- Etapas (método B) ----------------
    eb = calc["etapas_b"]
    etapas_tot = None
    if eb:
        ws = wb.create_sheet("Etapas")
        _titulo(ws, "PERCENTUAL EXECUTADO POR ETAPAS (PESOS)", dados, 8)
        for col, w in zip("ABCDEFGH", [40, 12, 14, 14, 16, 16, 16, 40]):
            ws.column_dimensions[col].width = w
        cub = eb["cub"] or {}
        _cel(ws, 5, 1, f"CUB/m² — {cub.get('projeto_padrao','')} — {cub.get('fonte','')} {cub.get('mes','')}".strip(" —"), F_BOLD)
        _cel(ws, 5, 2, cub.get("valor") if cub.get("valor") not in ("", None) else None, F_IN, MOEDA)
        _cel(ws, 6, 1, "Área equivalente (m²) — NBR 12721", F_BOLD)
        _cel(ws, 6, 2, eb["area"], F_IN, QTD)
        _cel(ws, 7, 1, "Valor de contrato (se a base for o contrato)", F_BOLD)
        _cel(ws, 7, 2, eb["contrato"], F_IN, MOEDA)
        _cel(ws, 8, 1, "Itens não cobertos pelo CUB (fundações especiais, elevadores etc.)", F_BOLD)
        _cel(ws, 8, 2, eb["adicional"], F_IN, MOEDA)
        _cel(ws, 9, 1, "VALOR GLOBAL DE REFERÊNCIA", F_BOLD, fill=FILL_TOT)
        if eb["base"] == "contrato":
            _cel(ws, 9, 2, "=N(B7)+N(B8)", F_BOLD, MOEDA, FILL_TOT)
        else:
            _cel(ws, 9, 2, "=N(B5)*N(B6)+N(B8)", F_BOLD, MOEDA, FILL_TOT)
        _cel(ws, 10, 1, f"Origem dos pesos: {eb['pesos_origem'] or 'não informada'}", F_SUB)
        HR = 12
        cab = ["Etapa", "Peso", "% executado (aceito)", "Contribuição", "Valor da etapa", "Valor executado", "Valor a executar", "Observações"]
        for i, h in enumerate(cab, start=1):
            c = ws.cell(row=HR, column=i, value=h)
            c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BORDA
        r0 = HR + 1
        for i, it in enumerate(eb["itens"]):
            r = r0 + i
            _cel(ws, r, 1, it["etapa"])
            _cel(ws, r, 2, it["peso"], F_IN, PCT)
            if it["pct"] is None:
                _cel(ws, r, 3, None, F_IN, PCT, FILL_PEND)
            else:
                _cel(ws, r, 3, it["pct"], F_IN, PCT)
            _cel(ws, r, 4, f"=B{r}*N(C{r})", fmt=PCT)
            _cel(ws, r, 5, f"=B{r}*$B$9", fmt=MOEDA)
            _cel(ws, r, 6, f"=E{r}*N(C{r})", fmt=MOEDA)
            _cel(ws, r, 7, f"=E{r}-F{r}", fmt=MOEDA)
            _cel(ws, r, 8, it["observacao"])
        rN = r0 + len(eb["itens"]) - 1
        etapas_tot = rN + 1
        _cel(ws, etapas_tot, 1, "TOTAL", F_BOLD, fill=FILL_TOT)
        for col, f in ((2, PCT), (4, PCT), (5, MOEDA), (6, MOEDA), (7, MOEDA)):
            L = get_column_letter(col)
            _cel(ws, etapas_tot, col, f"=SUM({L}{r0}:{L}{rN})", F_BOLD, f, FILL_TOT)
        _cel(ws, etapas_tot, 3, None, fill=FILL_TOT)
        _cel(ws, etapas_tot, 8, f'=IF(ABS(B{etapas_tot}-1)>0.0001,"ATENÇÃO: pesos não somam 100%","Pesos somam 100%")', F_BOLD, fill=FILL_TOT)

    # ---------------- Anomalias ----------------
    anomalias = dados.get("anomalias") or []
    if anomalias:
        ws = wb.create_sheet("Anomalias")
        _titulo(ws, "REGISTRO DE ANOMALIAS E NÃO CONFORMIDADES", dados, 11)
        cab = ["ID", "Local", "Sistema", "Anomalia", "Origem", "Natureza", "Prioridade", "Grau de risco", "Especificado × executado", "Normas / requisitos", "Recomendação", "Serviços relacionados", "Custo dos serviços relacionados"]
        _cab(ws, 5, cab, [7, 22, 16, 45, 20, 18, 10, 11, 30, 35, 50, 14, 16])
        for i, a in enumerate(anomalias):
            r = 6 + i
            _cel(ws, r, 1, a.get("id", ""), align=CENTER)
            _cel(ws, r, 2, a.get("local", ""))
            _cel(ws, r, 3, a.get("sistema", ""))
            _cel(ws, r, 4, a.get("titulo") or a.get("descricao", "")[:200])
            _cel(ws, r, 5, ORIGEM.get(a.get("origem"), a.get("origem", "")))
            _cel(ws, r, 6, NATUREZA.get(a.get("natureza"), a.get("natureza", "")))
            _cel(ws, r, 7, a.get("prioridade"), align=CENTER)
            _cel(ws, r, 8, GRAU_RISCO.get(a.get("_grau") or "", ""), align=CENTER)
            esp = " × ".join(x for x in [a.get("especificado", ""), a.get("executado", "")] if x)
            if a.get("referencia_projeto"):
                esp = f"{esp} ({a['referencia_projeto']})".strip()
            _cel(ws, r, 9, esp)
            _cel(ws, r, 10, "; ".join(a.get("normas") or []))
            _cel(ws, r, 11, a.get("recomendacao", ""))
            refs = [str(x) for x in (a.get("servicos") or [])]
            _cel(ws, r, 12, ", ".join(refs), align=CENTER)
            if refs and tot_row:
                partes = [f'SUMIF({rng("A")},"{x}",{rng("R")})' for x in refs]
                _cel(ws, r, 13, "=" + "+".join(partes), fmt=MOEDA, font=F_LINK)
            else:
                _cel(ws, r, 13, None, fmt=MOEDA)
        nota = 6 + len(anomalias) + 1
        ws.cell(row=nota, column=1, value="Nota: um mesmo serviço pode estar relacionado a mais de uma anomalia; não some a coluna de custo das anomalias.").font = F_SUB

    # ---------------- Materiais ----------------
    if calc["materiais"]:
        ws = wb.create_sheet("Materiais")
        _titulo(ws, "MATERIAIS ENTREGUES E NÃO INCORPORADOS À OBRA", dados, 7)
        _cab(ws, 5, ["Descrição", "Un", "Quantidade", "Estado de conservação", "Valor unitário", "Valor total", "Observações"], [45, 6, 12, 24, 14, 15, 40])
        for i, m in enumerate(calc["materiais"]):
            r = 6 + i
            _cel(ws, r, 1, m.get("descricao", ""))
            _cel(ws, r, 2, m.get("unidade", ""), align=CENTER)
            _cel(ws, r, 3, m["quantidade"], F_IN, QTD)
            _cel(ws, r, 4, m.get("estado", ""))
            _cel(ws, r, 5, m["valor_unitario"], F_IN, MOEDA, FILL_PEND if m["valor_unitario"] is None else None)
            _cel(ws, r, 6, f'=IF(E{r}="","",C{r}*E{r})', fmt=MOEDA)
            _cel(ws, r, 7, m.get("observacao", ""))
        rN = 5 + len(calc["materiais"])
        _cel(ws, rN + 1, 1, "TOTAL", F_BOLD, fill=FILL_TOT)
        _cel(ws, rN + 1, 6, f"=SUM(F6:F{rN})", F_BOLD, MOEDA, FILL_TOT)

    # ---------------- Premissas ----------------
    ws_p = wb.create_sheet("Premissas")
    _titulo(ws_p, "PREMISSAS, FONTES E LEGENDA", dados, 2)
    ws_p.column_dimensions["A"].width = 4
    ws_p.column_dimensions["B"].width = 120
    itens = []
    if base:
        itens.append(
            f"Base de custos: {base.get('fonte','')} — UF {base.get('uf','')} — mês de referência {base.get('mes_referencia','')} — "
            f"{'desonerado' if base.get('regime') == 'desonerado' else 'não desonerado'} — BDI {calc['bdi']*100:.2f}%"
            + (f" — consulta em {base.get('data_consulta')}" if base.get("data_consulta") else "")
        )
    for o in base.get("observacoes") or []:
        itens.append(o)
    if eb:
        itens.append(f"Pesos por etapa: {eb['pesos_origem'] or 'origem não informada'}.")
    itens += [
        "Qtd. não executada = prevista − executada aceita − a refazer (mínimo zero).",
        "Custo de refazer = qtd. a refazer × (PU do serviço + PU de demolição/remoção) × (1 + BDI).",
        "% executado aceito = Σ valor executado aceito ÷ Σ valor previsto (serviços com preço). Serviço executado com vício não é considerado executado.",
        "Serviços sem preço unitário ficam fora dos totais até o preenchimento.",
        "LEGENDA: texto azul = dado de entrada; preto = fórmula; verde = vínculo com outra aba; fundo amarelo = preenchimento pendente.",
    ]
    for i, t in enumerate(itens):
        c = ws_p.cell(row=5 + i, column=2, value=t)
        c.font = F_TXT
        c.alignment = WRAP

    # ---------------- Resumo ----------------
    ws = ws_res
    _titulo(ws, f"RESUMO — {TIPOS.get(tipo, {}).get('titulo', 'LAUDO')}", dados, 6)
    ws.column_dimensions["A"].width = 44
    for col in "BCDEF":
        ws.column_dimensions[col].width = 18
    r = 5
    if tot_row:
        _cel(ws, r, 1, "ORÇAMENTO DISCRIMINADO", F_HEAD, fill=FILL_HEAD)
        _cel(ws, r, 2, "Valor", F_HEAD, fill=FILL_HEAD, align=CENTER)
        r += 1
        pares = [
            ("Valor total previsto", f"='Orçamento'!N{tot_row}", MOEDA),
            ("Valor executado aceito", f"='Orçamento'!O{tot_row}", MOEDA),
            ("% executado aceito", f"=IF('Orçamento'!N{tot_row}=0,\"\",'Orçamento'!O{tot_row}/'Orçamento'!N{tot_row})", PCT),
            ("% executado bruto (inclui serviços com vício)", f"=IF('Orçamento'!N{tot_row}=0,\"\",'Orçamento'!S{tot_row}/'Orçamento'!N{tot_row})", PCT),
            ("Custo de refazer (vícios)", f"='Orçamento'!P{tot_row}", MOEDA),
            ("Custo a executar (não executado / pendências)", f"='Orçamento'!Q{tot_row}", MOEDA),
            ("CUSTO TOTAL PARA CONCLUIR / CORRIGIR", f"='Orçamento'!R{tot_row}", MOEDA),
            ("% físico informado (programa / medições)", _pct_inf(dados), PCT),
            ("Serviços sem preço (pendentes)", f"=COUNTBLANK('Orçamento'!K{tot_row - len(linhas_s)}:K{tot_row - 1})", "0"),
        ]
        for k, f, fmt in pares:
            bold = k.isupper()
            _cel(ws, r, 1, k, F_BOLD if bold else F_TXT, fill=FILL_TOT if bold else None)
            _cel(ws, r, 2, f, F_BOLD if bold else F_LINK, fmt, FILL_TOT if bold else None)
            r += 1
        r += 1
        # por etapa
        cab = ["Etapa", "Valor previsto", "Valor aceito", "% aceito", "Custo de refazer", "Custo a executar"]
        for i, h in enumerate(cab, start=1):
            _cel(ws, r, i, h, F_HEAD, fill=FILL_HEAD, align=CENTER)
        r += 1
        for e in calc["por_etapa"]:
            ee = e["etapa"].replace('"', '""')
            _cel(ws, r, 1, e["etapa"])
            _cel(ws, r, 2, f'=SUMIF({rng("B")},"{ee}",{rng("N")})', F_LINK, MOEDA)
            _cel(ws, r, 3, f'=SUMIF({rng("B")},"{ee}",{rng("O")})', F_LINK, MOEDA)
            _cel(ws, r, 4, f'=IF(B{r}=0,"",C{r}/B{r})', fmt=PCT)
            _cel(ws, r, 5, f'=SUMIF({rng("B")},"{ee}",{rng("P")})', F_LINK, MOEDA)
            _cel(ws, r, 6, f'=SUMIF({rng("B")},"{ee}",{rng("Q")})', F_LINK, MOEDA)
            r += 1
        r += 1
    if etapas_tot:
        _cel(ws, r, 1, "MÉTODO POR ETAPAS (PESOS)", F_HEAD, fill=FILL_HEAD)
        _cel(ws, r, 2, "Valor", F_HEAD, fill=FILL_HEAD, align=CENTER)
        r += 1
        for k, f, fmt in [
            ("Valor global de referência", "=Etapas!B9", MOEDA),
            ("% executado (ponderado)", f"=Etapas!D{etapas_tot}", PCT),
            ("Valor executado", f"=Etapas!F{etapas_tot}", MOEDA),
            ("Valor a executar", f"=Etapas!G{etapas_tot}", MOEDA),
        ]:
            _cel(ws, r, 1, k)
            _cel(ws, r, 2, f, F_LINK, fmt)
            r += 1
        r += 1
    if anomalias:
        _cel(ws, r, 1, "ANOMALIAS POR PRIORIDADE", F_HEAD, fill=FILL_HEAD)
        _cel(ws, r, 2, "Quantidade", F_HEAD, fill=FILL_HEAD, align=CENTER)
        r += 1
        n = 5 + len(anomalias)
        for p in (1, 2, 3):
            _cel(ws, r, 1, f"Prioridade {p}")
            _cel(ws, r, 2, f"=COUNTIF(Anomalias!G6:G{n},{p})", F_LINK, "0")
            r += 1
        _cel(ws, r, 1, "Total de anomalias", F_BOLD)
        _cel(ws, r, 2, f"=COUNTA(Anomalias!A6:A{n})", F_BOLD, "0")
        r += 1

    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth = 1
        w.page_setup.fitToHeight = 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(caminho)
