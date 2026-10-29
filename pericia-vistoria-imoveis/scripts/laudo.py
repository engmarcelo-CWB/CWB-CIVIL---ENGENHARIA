"""Gera o laudo .docx (python-docx)."""

import os
import re
import tempfile
from datetime import date, datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from textos_padrao import (
    GRAU_RISCO,
    NATUREZA,
    ORIGEM,
    PRIORIDADE,
    SISTEMAS,
    SITUACAO_CHECK,
    SITUACAO_DOC,
    SITUACAO_SISTEMA,
    TIPOS,
)

FONTE = "Arial"
AZUL = RGBColor(0x1F, 0x38, 0x64)
CINZA = RGBColor(0x59, 0x59, 0x59)
COR_CAB = "1F3864"
COR_ZEBRA = "EEF2F8"
COR_TOT = "D9E1F2"


# ------------------------------------------------------------------ formatação
def moeda(v):
    if v is None:
        return "a definir"
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def num(v, d=2):
    if v is None:
        return "—"
    s = f"{v:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return s


def pct(v):
    if v is None:
        return "—"
    return num(v * 100, 2) + "%"


def data_br(s):
    if not s:
        return ""
    if isinstance(s, (date, datetime)):
        return s.strftime("%d/%m/%Y")
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", str(s))
    return f"{m.group(3)}/{m.group(2)}/{m.group(1)}" if m else str(s)


MESES = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]


def data_extenso(s):
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", str(s or ""))
    if not m:
        return str(s or "")
    return f"{int(m.group(3))} de {MESES[int(m.group(2)) - 1]} de {m.group(1)}"


# ------------------------------------------------------------------ helpers docx
def _shade(cell, cor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcPr.append(shd)


def _campo(run, instr):
    for tag, text in (("begin", None), (None, instr), ("separate", None), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {text} "
            run._r.append(it)


def _larguras(t, larguras):
    """Fixa larguras de coluna (Word e LibreOffice): layout fixo + gridCol + largura de cada célula."""
    t.autofit = False
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for i, w in enumerate(larguras):
        t.columns[i].width = Cm(w)
        for cell in t.columns[i].cells:
            cell.width = Cm(w)


def _repetir_cabecalho(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _nao_quebrar(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _exif_meta(img):
    """Data/hora e coordenadas GPS do EXIF, quando existirem."""
    meta = {}
    try:
        ex = img.getexif()
        sub = ex.get_ifd(0x8769) if hasattr(ex, "get_ifd") else {}
        dt = sub.get(36867) or ex.get(306)
        if dt:
            m = re.match(r"(\d{4}):(\d{2}):(\d{2})[ T](\d{2}):(\d{2})", str(dt))
            if m:
                meta["data"] = f"{m.group(3)}/{m.group(2)}/{m.group(1)} {m.group(4)}:{m.group(5)}"
        gps = ex.get_ifd(0x8825) if hasattr(ex, "get_ifd") else {}
        if gps and 2 in gps and 4 in gps:
            def dec(v, ref):
                g = float(v[0]) + float(v[1]) / 60 + float(v[2]) / 3600
                return -g if ref in ("S", "W") else g
            lat, lon = dec(gps[2], gps.get(1, "N")), dec(gps[4], gps.get(3, "E"))
            meta["coord"] = f"{lat:.5f}, {lon:.5f}"
    except Exception:  # noqa: BLE001
        pass
    return meta


class Laudo:
    def __init__(self, dados, calc, base_dir, perito):
        self.d = dados
        self.c = calc
        self.base_dir = base_dir
        self.perito = perito
        self.tipo = dados.get("tipo", "")
        self.t = TIPOS.get(self.tipo, TIPOS["patologias"])
        self.doc = Document()
        self.sec = 0
        self.sub = 0
        self.n_foto = 0
        self.n_tab = 0
        self.tmp = tempfile.mkdtemp(prefix="laudo_img_")
        self.avisos = []
        self.toc = []
        self._toc_run = None
        self._estilos()

    # ---------------------------------------------------------- estilos
    def _estilos(self):
        st = self.doc.styles
        n = st["Normal"]
        n.font.name = FONTE
        n.font.size = Pt(11)
        n.element.rPr.rFonts.set(qn("w:eastAsia"), FONTE)
        n.paragraph_format.space_after = Pt(6)
        n.paragraph_format.line_spacing = 1.15
        for nome, tam in (("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 11), ("Heading 4", 10.5)):
            h = st[nome]
            h.font.name = FONTE
            h.font.size = Pt(tam)
            h.font.bold = True
            h.font.color.rgb = AZUL
            rf = h.element.rPr.rFonts
            for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rf.get(qn(att)) is not None:
                    del rf.attrib[qn(att)]
            for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rf.set(qn(att), FONTE)
            h.paragraph_format.space_before = Pt(14 if nome == "Heading 1" else 10)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        for nome in ("List Bullet",):
            s = st[nome]
            s.font.name = FONTE
            s.font.size = Pt(11)
        sec = self.doc.sections[0]
        sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
        sec.top_margin, sec.left_margin = Cm(2.5), Cm(3.0)
        sec.bottom_margin, sec.right_margin = Cm(2.0), Cm(2.0)
        sec.header_distance = Cm(1.0)
        sec.footer_distance = Cm(1.0)

    # ---------------------------------------------------------- blocos
    def h1(self, texto):
        self.sec += 1
        self.sub = 0
        titulo = f"{self.sec}. {texto.upper()}"
        self.toc.append((1, titulo))
        self.doc.add_heading(titulo, level=1)

    def h2(self, texto):
        self.sub += 1
        titulo = f"{self.sec}.{self.sub}. {texto}"
        self.toc.append((2, titulo))
        self.doc.add_heading(titulo, level=2)

    def p(self, texto, bold=False, italic=False, align="just", size=None, cor=None, after=None):
        par = self.doc.add_paragraph()
        if align == "just":
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "right":
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._runs(par, texto, bold, italic, size, cor)
        if after is not None:
            par.paragraph_format.space_after = Pt(after)
        return par

    def _runs(self, par, texto, bold=False, italic=False, size=None, cor=None):
        # suporta **negrito** inline
        partes = re.split(r"(\*\*[^*]+\*\*)", str(texto))
        for pt in partes:
            if not pt:
                continue
            b = bold
            if pt.startswith("**") and pt.endswith("**"):
                pt, b = pt[2:-2], True
            r = par.add_run(pt)
            r.bold = b
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if cor:
                r.font.color.rgb = cor

    def texto(self, texto):
        """Parágrafos separados por linha em branco; linhas iniciadas por '- ' viram marcadores."""
        if not texto:
            return
        if isinstance(texto, list):
            for t in texto:
                self.texto(t)
            return
        for bloco in re.split(r"\n\s*\n", str(texto).strip()):
            linhas = bloco.split("\n")
            if all(l.strip().startswith(("- ", "• ")) for l in linhas if l.strip()):
                for l in linhas:
                    if l.strip():
                        self.bullet(l.strip()[2:])
            else:
                self.p(" ".join(l.strip() for l in linhas))

    def bullet(self, texto):
        par = self.doc.add_paragraph(style="List Bullet")
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.paragraph_format.space_after = Pt(3)
        self._runs(par, texto)
        return par

    def legenda_tabela(self, texto):
        self.n_tab += 1
        par = self.p(f"Tabela {self.n_tab} – {texto}", bold=True, align="center", size=9, after=2)
        par.paragraph_format.keep_with_next = True

    def tabela(self, cab, linhas, larguras=None, tam=9, total=None, alinhamentos=None):
        t = self.doc.add_table(rows=1, cols=len(cab))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = t.rows[0]
        _repetir_cabecalho(hdr)
        for i, h in enumerate(cab):
            cell = hdr.cells[i]
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = par.add_run(str(h))
            r.bold = True
            r.font.size = Pt(tam)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell, COR_CAB)
        todas = list(linhas) + ([total] if total else [])
        for j, ln in enumerate(todas):
            row = t.add_row()
            _nao_quebrar(row)
            eh_total = total is not None and j == len(todas) - 1
            for i, v in enumerate(ln):
                cell = row.cells[i]
                cell.text = ""
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(0)
                if alinhamentos and alinhamentos[i] == "r":
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                elif alinhamentos and alinhamentos[i] == "c":
                    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                self._runs(par, "" if v is None else str(v), bold=eh_total, size=tam)
                if eh_total:
                    _shade(cell, COR_TOT)
                elif j % 2 == 1:
                    _shade(cell, COR_ZEBRA)
        _larguras(t, larguras or [16.0 / len(cab)] * len(cab))
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def tabela_chave_valor(self, pares, larg=(5.0, 11.0)):
        pares = [(k, v) for k, v in pares if v not in (None, "", [])]
        if not pares:
            return
        t = self.doc.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, v in pares:
            row = t.add_row()
            _nao_quebrar(row)
            c0, c1 = row.cells
            c0.text, c1.text = "", ""
            r = c0.paragraphs[0].add_run(k)
            r.bold = True
            r.font.size = Pt(10)
            _shade(c0, COR_ZEBRA)
            self._runs(c1.paragraphs[0], str(v), size=10)
        _larguras(t, larg)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---------------------------------------------------------- fotos
    def _preparar_imagem(self, arquivo):
        caminho = arquivo if os.path.isabs(arquivo) else os.path.join(self.base_dir, arquivo)
        if not os.path.exists(caminho):
            self.avisos.append(f"Foto não encontrada: {arquivo}")
            return None
        try:
            from PIL import Image, ImageOps

            img = Image.open(caminho)
            meta = _exif_meta(img)
            img = ImageOps.exif_transpose(img)
            if img.mode not in ("RGB", "L"):
                img = img.convert("RGB")
            img.thumbnail((1600, 1600))
            saida = os.path.join(self.tmp, f"foto_{self.n_foto + 1:03d}.jpg")
            img.save(saida, "JPEG", quality=82)
            return saida, img.size, meta
        except Exception as e:  # noqa: BLE001
            self.avisos.append(f"Não foi possível abrir a foto {arquivo}: {e}")
            return None

    def fotos(self, lista, prefixo=""):
        lista = [f for f in (lista or []) if f and f.get("arquivo")]
        if not lista:
            return []
        numeros = []
        t = self.doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i in range(0, len(lista), 2):
            row = t.add_row()
            _nao_quebrar(row)
            for j in range(2):
                cell = row.cells[j]
                cell.width = Cm(8.0)
                if i + j >= len(lista):
                    continue
                f = lista[i + j]
                prep = self._preparar_imagem(f["arquivo"])
                self.n_foto += 1
                numeros.append(self.n_foto)
                par = cell.paragraphs[0]
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if prep:
                    caminho, (w, h), meta = prep
                    # paisagem: 7,6 cm de largura; retrato: limita a altura a ~8,7 cm
                    largura = 7.6 if w >= h else max(5.0, min(7.6, 8.7 * w / h))
                    par.add_run().add_picture(caminho, width=Cm(largura))
                else:
                    self._runs(par, f"[foto não encontrada: {f['arquivo']}]", italic=True, size=9)
                cap = cell.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(6)
                leg = f.get("legenda", "")
                extra = []
                data_f = f.get("data") or (meta.get("data") if prep and self.d.get("fotos_exif", True) else None)
                coord_f = f.get("coordenadas") or (meta.get("coord") if prep and self.d.get("fotos_exif", True) else None)
                if data_f:
                    extra.append(data_br(data_f))
                if coord_f:
                    extra.append(str(coord_f))
                txt = f"Foto {self.n_foto}{(' – ' + leg) if leg else ''}"
                if extra:
                    txt += f" ({'; '.join(extra)})"
                if f.get("fonte"):
                    txt += f". Fonte: {f['fonte']}"
                self._runs(cap, txt, italic=True, size=9)
        _larguras(t, [8.0, 8.0])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return numeros

    # ---------------------------------------------------------- cabeçalho/rodapé/capa
    def _cabecalho_rodape(self, sec):
        emp = self.perito.get("empresa") or {}
        hdr = sec.header
        hdr.is_linked_to_previous = False
        par = hdr.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        logo = self.perito.get("logo")
        if logo and os.path.exists(logo):
            par.add_run().add_picture(logo, height=Cm(1.0))
            par.add_run("\n")
        self._runs(par, emp.get("razao_social", self.perito.get("nome", "")), bold=True, size=8, cor=AZUL)
        linha2 = " | ".join(x for x in [f"CNPJ {emp['cnpj']}" if emp.get("cnpj") else "", f"Laudo nº {self.d.get('numero_laudo','')}"] if x)
        p2 = hdr.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._runs(p2, linha2, size=8, cor=CINZA)
        ftr = sec.footer
        ftr.is_linked_to_previous = False
        fp = ftr.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run(" | ".join(x for x in [emp.get("endereco", ""), emp.get("telefone", ""), emp.get("email", "")] if x))
        r.font.size = Pt(7)
        r.font.color.rgb = CINZA
        fp2 = ftr.add_paragraph()
        fp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for txt, campo in (("Página ", "PAGE"), (" de ", "NUMPAGES")):
            rr = fp2.add_run(txt)
            rr.font.size = Pt(8)
            rc = fp2.add_run()
            rc.font.size = Pt(8)
            _campo(rc, campo)

    def capa(self):
        d = self.d
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        self._cabecalho_rodape(sec)
        emp = self.perito.get("empresa") or {}
        im = d.get("imovel") or {}
        for _ in range(2):
            self.doc.add_paragraph()
        self.p(emp.get("razao_social", ""), bold=True, align="center", size=12, cor=AZUL)
        for _ in range(3):
            self.doc.add_paragraph()
        self.p(d.get("titulo") or self.t["titulo"], bold=True, align="center", size=18, cor=AZUL)
        if d.get("subtitulo"):
            self.p(d["subtitulo"], bold=True, align="center", size=12)
        self.p(f"Laudo nº {d.get('numero_laudo','')}", align="center", size=12)
        self.p(self.t["especie"], italic=True, align="center", size=10, cor=CINZA)
        self.doc.add_paragraph()
        # bloco de apresentação do objeto (padrão dos laudos de referência)
        objeto = d.get("objeto") or f"imóvel situado em {im.get('endereco','')}{', ' + im['municipio_uf'] if im.get('municipio_uf') else ''}"
        sol = d.get("solicitante") or {}
        bloco = self.doc.add_paragraph()
        bloco.paragraph_format.left_indent = Cm(7.5)
        bloco.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        self._runs(
            bloco,
            f"Laudo técnico elaborado por profissional legalmente habilitado, por solicitação de {sol.get('nome','')}, "
            f"tendo por objeto: {objeto.rstrip('.')}. {('Finalidade: ' + d['finalidade']) if d.get('finalidade') else ''}".strip(),
            size=10,
        )
        self.doc.add_paragraph()
        proc = d.get("processo") or {}
        prog = d.get("programa") or {}
        pares = [
            ("Imóvel", ", ".join(x for x in [im.get("endereco", ""), im.get("municipio_uf", "")] if x)),
            ("Solicitante", sol.get("nome", "")),
        ]
        if proc.get("numero"):
            pares.append(("Processo", f"{proc.get('numero')} — {proc.get('juizo','')} {('— ' + proc['comarca']) if proc.get('comarca') else ''}".strip(" —")))
        if prog.get("id_sistema") or prog.get("instrumento"):
            pares.append(("Obra / instrumento", " — ".join(x for x in [prog.get("instrumento", ""), prog.get("id_sistema", "")] if x)))
        pares.append(("Responsável técnico", f"{self.perito.get('titulo','')} {self.perito.get('nome','')} — {self.perito.get('registro','')}".strip()))
        pares.append(("ART", self.perito.get("art") or "a registrar"))
        self.tabela_chave_valor(pares, larg=(4.0, 12.0))
        for _ in range(2):
            self.doc.add_paragraph()
        self.p(f"{d.get('local_emissao','')}, {data_extenso(d.get('data_emissao',''))}".strip(", "), align="center", size=11)
        self.doc.add_page_break()

    def sumario(self):
        par = self.doc.add_paragraph()
        self._runs(par, "SUMÁRIO", bold=True, size=13, cor=AZUL)
        par2 = self.doc.add_paragraph()
        for tag, texto in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None)):
            r = par2.add_run()
            if tag:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), tag)
                r._r.append(el)
            else:
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = f" {texto} "
                r._r.append(it)
        self._toc_run = par2.add_run()  # recebe a lista de seções ao final (visível mesmo sem atualizar o campo)
        r = par2.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), "end")
        r._r.append(el)
        nota = self.doc.add_paragraph()
        self._runs(nota, "(No Word, o sumário é atualizado com os números de página ao abrir o arquivo; se não for, clique com o botão direito sobre ele e escolha \"Atualizar campo\".)", italic=True, size=8, cor=CINZA)
        self.doc.add_page_break()
        # pede ao Word para atualizar os campos ao abrir
        settings = self.doc.settings.element
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)

    # ---------------------------------------------------------- seções
    def identificacao(self):
        d = self.d
        self.h1("Identificação")
        sol = d.get("solicitante") or {}
        self.h2("Solicitante")
        self.tabela_chave_valor([
            ("Nome / razão social", sol.get("nome")),
            ("CPF / CNPJ", sol.get("documento")),
            ("Qualidade", sol.get("qualidade")),
            ("Endereço", sol.get("endereco")),
            ("Contato", sol.get("contato")),
        ])
        prop = d.get("proprietario") or {}
        if prop.get("nome"):
            self.h2("Proprietário")
            self.tabela_chave_valor([("Nome", prop.get("nome")), ("CPF / CNPJ", prop.get("documento")), ("Endereço", prop.get("endereco"))])
        proc = d.get("processo") or {}
        if proc.get("numero"):
            self.h2("Processo")
            self.tabela_chave_valor([
                ("Número", proc.get("numero")), ("Juízo", proc.get("juizo")), ("Comarca", proc.get("comarca")),
                ("Partes", proc.get("partes")),
            ])
        self.h2("Responsável técnico")
        emp = self.perito.get("empresa") or {}
        regs = ", ".join([self.perito.get("registro", "")] + list(self.perito.get("registros_adicionais") or []))
        self.tabela_chave_valor([
            ("Nome", self.perito.get("nome")),
            ("Título", self.perito.get("titulo")),
            ("Registro profissional", regs),
            ("ART do laudo", self.perito.get("art") or "a registrar"),
            ("Empresa", f"{emp.get('razao_social','')} — CNPJ {emp.get('cnpj','')}" if emp.get("razao_social") else None),
            ("Contato", " | ".join(x for x in [emp.get("telefone", ""), emp.get("email", "")] if x)),
        ])

    def objeto(self):
        d = self.d
        self.h1("Objeto, objetivo e finalidade")
        im = d.get("imovel") or {}
        objeto = d.get("objeto") or f"Imóvel situado em {im.get('endereco','')}{', ' + im['municipio_uf'] if im.get('municipio_uf') else ''}."
        self.p(f"**Objeto:** {objeto}")
        self.p(f"**Espécie:** {self.t['especie']}.")
        self.p(f"**Objetivo:** {d.get('objetivo') or self.t['objetivo']}")
        if d.get("finalidade"):
            self.p(f"**Finalidade:** {d['finalidade']}")

    def dados_obra(self):
        prog = self.d.get("programa") or {}
        if not prog:
            return
        self.h1("Dados da obra")
        self.tabela_chave_valor([
            ("Programa / órgão", prog.get("orgao")),
            ("Norma do programa", prog.get("norma_referencia")),
            ("Instrumento", prog.get("instrumento")),
            ("Identificação no sistema", prog.get("id_sistema")),
            ("Objeto pactuado", prog.get("objeto")),
            ("Empresa executora", prog.get("executor")),
            ("Contrato", prog.get("contrato")),
            ("Valor pactuado", moeda(float(prog["valor_pactuado"])) if isinstance(prog.get("valor_pactuado"), (int, float)) else prog.get("valor_pactuado")),
            ("Prazo de execução", prog.get("prazo")),
            ("Início da execução", data_br(prog.get("data_inicio"))),
            ("Paralisação", data_br(prog.get("data_paralisacao"))),
            ("% físico informado", f"{pct(_frac(prog['percentual_informado']))} — fonte: {prog.get('fonte_percentual','não informada')}" if prog.get("percentual_informado") not in (None, "") else None),
            ("Situação no sistema", prog.get("situacao")),
        ])
        if prog.get("historico"):
            self.texto(prog["historico"])
        if prog.get("imagem_sistema"):
            self.fotos([prog["imagem_sistema"]])

    def pressupostos(self):
        self.h1("Pressupostos, ressalvas e condicionantes")
        itens = list(self.t["pressupostos"]) + list(self.d.get("pressupostos") or [])
        itens.append(
            "O signatário declara não ter interesse no bem vistoriado nem relação com as partes que comprometa sua isenção, "
            "e que as análises e conclusões refletem seu entendimento técnico, observado o Código de Ética Profissional do Sistema Confea/Crea."
        )
        for i in itens:
            self.bullet(i)

    def normas(self):
        self.h1("Referências normativas e técnicas")
        lista, vistos = [], set()

        def add(n, t=""):
            m = re.search(r"NBR\s*(ISO\s*)?\d+(-\d+)?", n)
            k = m.group(0).replace(" ", "").upper() if m else n.strip().lower()
            if k and k not in vistos:
                vistos.add(k)
                lista.append((n, t))

        for n, t in self.t["normas"]:
            add(n, t)
        sist = (self.d.get("imovel") or {}).get("sistema_construtivo_tipo")
        for n, t in SISTEMAS.get(sist, ("", []))[1]:
            add(n, t)
        for n in self.d.get("normas") or []:
            if isinstance(n, dict):
                add(n.get("norma", ""), n.get("titulo", ""))
            else:
                add(str(n))
        for a in self.d.get("anomalias") or []:
            for n in a.get("normas") or []:
                add(str(n))
        for n, t in lista:
            self.bullet(f"{n}{' — ' + t if t else ''}")

    def metodologia(self):
        self.h1("Metodologia e critérios")
        self.h2("Procedimentos")
        met = self.d.get("metodologia") or self.t["metodologia"]
        for i, m in enumerate(met, 1):
            self.p(f"{i}) {m}", after=3)
        if self.d.get("nivel_inspecao"):
            self.p(f"**Nível de inspeção:** {self.d['nivel_inspecao']}")
        if self.d.get("anomalias"):
            self.h2("Critérios de classificação")
            self.p(
                "As anomalias e não conformidades foram classificadas quanto à **origem** (endógena — projeto, material ou execução; "
                "exógena; funcional; uso, operação ou manutenção) e à **natureza** (vício aparente ou oculto, defeito, avaria, mutilação, "
                "deterioração, serviço não executado, não conformidade), com os conceitos da ABNT NBR 13752:2024, e organizadas em "
                "patamares de prioridade conforme a Norma de Inspeção Predial do IBAPE (2025), com a correspondência ao grau de risco:"
            )
            self.tabela(
                ["Prioridade", "Grau de risco", "Critério"],
                [
                    ("1", "Crítico", "Compromete saúde, segurança ou funcionalidade com risco de paralisação; perda de vida útil, aumento expressivo de custo ou risco ambiental. Ação imediata."),
                    ("2", "Regular", "Funcionalidade parcialmente comprometida, sem prejuízo à operação e sem afetar saúde ou segurança; deterioração precoce."),
                    ("3", "Mínimo", "Pequenas perdas de desempenho, estéticas ou programáveis, sem comprometer vida útil ou durabilidade."),
                ],
                larguras=[2.2, 2.6, 11.2], tam=8, alinhamentos=["c", "c", "l"],
            )
            if self.d.get("criterio_fissuras"):
                self.p(f"**Fissuras:** {self.d['criterio_fissuras']}")

    def documentos(self):
        docs = self.d.get("documentos") or []
        self.h1("Documentação analisada")
        if not docs:
            self.p("Não foram disponibilizados documentos técnicos para análise. As constatações baseiam-se exclusivamente na vistoria e nas informações prestadas pelo solicitante.")
            return
        linhas = [(x.get("documento", ""), SITUACAO_DOC.get(x.get("situacao"), x.get("situacao", "")), x.get("observacao", "")) for x in docs]
        self.legenda_tabela("Documentos solicitados e analisados")
        self.tabela(["Documento", "Situação", "Observações"], linhas, larguras=[6.0, 3.8, 6.2])
        if self.d.get("analise_documental"):
            self.texto(self.d["analise_documental"])

    def caracterizacao(self):
        im = self.d.get("imovel") or {}
        self.h1("Caracterização e localização")

        def area(v):
            return f"{num(float(v))} m²" if isinstance(v, (int, float)) else (v or None)

        sist_cod = im.get("sistema_construtivo_tipo")
        sist_txt = im.get("sistema_construtivo") or SISTEMAS.get(sist_cod, ("", []))[0]
        self.tabela_chave_valor([
            ("Endereço", im.get("endereco")),
            ("Município/UF", im.get("municipio_uf")),
            ("Matrícula", im.get("matricula")),
            ("Inscrição imobiliária", im.get("inscricao")),
            ("Coordenadas", im.get("coordenadas")),
            ("Tipologia / uso", im.get("tipologia")),
            ("Padrão construtivo", im.get("padrao")),
            ("Sistema construtivo", sist_txt),
            ("Fabricante / sistemista", im.get("fabricante_sistema")),
            ("Pavimentos", im.get("pavimentos")),
            ("Área do terreno", area(im.get("area_terreno"))),
            ("Área construída", area(im.get("area_construida"))),
            ("Área privativa", area(im.get("area_privativa"))),
            ("Área equivalente (NBR 12721)", area(im.get("area_equivalente"))),
            ("Idade / datas de referência", im.get("idade")),
            ("Executor / construtora", im.get("executor")),
            ("Responsável técnico pela execução", im.get("responsavel_tecnico_execucao")),
            ("Serviços públicos disponíveis", im.get("servicos_publicos")),
        ])
        if im.get("imagem_localizacao"):
            self.fotos([im["imagem_localizacao"]])
        if im.get("descricao"):
            self.h2("Descrição técnica")
            self.texto(im["descricao"])
        carac = im.get("caracteristicas") or []
        if carac:
            self.h2("Características construtivas")
            self.legenda_tabela("Características construtivas por sistema")
            self.tabela(["Sistema / elemento", "Descrição"], [(c.get("item", ""), c.get("descricao", "")) for c in carac], larguras=[5.0, 11.0])

    def vistoria(self):
        v = self.d.get("vistoria") or {}
        self.h1("Vistoria")
        datas = ", ".join(data_br(x) for x in v.get("datas") or [])
        self.tabela_chave_valor([
            ("Data(s)", datas),
            ("Condições", v.get("condicoes")),
            ("Acompanhantes", ", ".join(v.get("acompanhantes") or [])),
            ("Instrumentos", ", ".join(v.get("instrumentos") or [])),
        ])
        if v.get("informacoes_colhidas"):
            self.h2("Informações colhidas (anamnese)")
            self.texto(v["informacoes_colhidas"])
        si = v.get("sistemas_inspecionados") or []
        if si:
            self.h2("Sistemas e áreas vistoriados")
            self.p("; ".join(si) + ".")
        sni = v.get("sistemas_nao_inspecionados") or []
        self.h2("Sistemas e áreas não vistoriados")
        if sni:
            self.legenda_tabela("Sistemas não vistoriados e motivos")
            self.tabela(["Sistema / área", "Motivo"], [(x.get("sistema", ""), x.get("motivo", "")) for x in sni], larguras=[6.0, 10.0])
        else:
            self.p("Todos os sistemas e áreas acessíveis foram vistoriados.")
        if self.d.get("fotos_gerais"):
            self.p("O registro fotográfico geral consta do Anexo — Relatório fotográfico.")

    def checklist(self):
        ck = self.d.get("checklist") or []
        if not ck:
            return
        self.h2("Verificação por ambiente e sistema")
        self.legenda_tabela("Verificação de conformidade por ambiente/sistema")
        linhas = [(c.get("ambiente", ""), c.get("sistema", ""), c.get("item", ""), SITUACAO_CHECK.get(c.get("situacao"), c.get("situacao", "")), c.get("anomalia", "")) for c in ck]
        self.tabela(["Ambiente", "Sistema", "Item verificado", "Situação", "Anomalia"], linhas, larguras=[3.0, 2.8, 6.0, 2.4, 1.8], tam=8, alinhamentos=["l", "l", "l", "c", "c"])

    def _anomalia(self, a, nivel=3):
        self.doc.add_heading(f"{a.get('id','')} — {a.get('titulo') or a.get('local','')}", level=nivel)
        grau = GRAU_RISCO.get(a.get("_grau") or "", "")
        self.tabela_chave_valor([
            ("Local", a.get("local")),
            ("Sistema", a.get("sistema")),
            ("Especificado", a.get("especificado")),
            ("Executado / medido", a.get("executado")),
            ("Referência de projeto", a.get("referencia_projeto")),
            ("Abertura (fissuras)", f"{num(float(a['abertura_mm']))} mm" if isinstance(a.get("abertura_mm"), (int, float)) else a.get("abertura_mm")),
            ("Origem", ORIGEM.get(a.get("origem"), a.get("origem"))),
            ("Natureza", NATUREZA.get(a.get("natureza"), a.get("natureza"))),
            ("Prioridade / grau de risco", f"{PRIORIDADE.get(a.get('prioridade'), a.get('prioridade') or '')}{' — grau ' + grau if grau else ''}" if (a.get("prioridade") or grau) else None),
        ], larg=(3.8, 12.2))
        self.p("**Descrição:**", after=2)
        self.texto(a.get("descricao", ""))
        if a.get("causa"):
            self.p("**Mecanismo e causa:**", after=2)
            self.texto(a["causa"])
        if a.get("nexo_causal"):
            self.p("**Nexo causal:**", after=2)
            self.texto(a["nexo_causal"])
        if a.get("normas"):
            self.p("**Requisitos normativos / de projeto:**", after=2)
            for n in a["normas"]:
                self.bullet(n)
        if a.get("recomendacao"):
            self.p("**Correção recomendada:**", after=2)
            self.texto(a["recomendacao"])
        if a.get("ensaios"):
            self.p("**Investigações complementares recomendadas:**", after=2)
            for e in a["ensaios"]:
                self.bullet(e)
        custo = self.c["custo_anomalia"].get(a.get("id"))
        if a.get("servicos"):
            self.p(f"**Serviços orçados:** itens {', '.join(str(s) for s in a['servicos'])} da planilha — custo estimado {moeda(custo)}.")
        self.fotos(a.get("fotos"))

    def constatacoes(self):
        anom = self.d.get("anomalias") or []
        sits = self.d.get("situacao_sistemas") or []
        titulo = {
            "obra_inacabada": "Constatações: estágio da obra e não conformidades",
            "patologias": "Constatações: manifestações patológicas",
            "inspecao_predial": "Constatações: anomalias e falhas",
            "recebimento_obra": "Constatações: pendências e não conformidades",
        }.get(self.tipo, "Constatações")
        self.h1(titulo)
        if self.d.get("constatacoes_gerais"):
            self.texto(self.d["constatacoes_gerais"])
        if sits:
            self.h2("Situação por sistema")
            self.legenda_tabela("Situação de cada sistema construtivo na data da vistoria")
            self.tabela(
                ["Sistema", "Situação", "Descrição"],
                [(s.get("sistema", ""), SITUACAO_SISTEMA.get(s.get("situacao"), s.get("situacao", "")), s.get("descricao", "")) for s in sits],
                larguras=[3.6, 3.0, 9.4], tam=8,
            )
        if self.tipo == "recebimento_obra":
            self.checklist()
        restr = (self.d.get("programa") or {}).get("restricoes") or []
        if restr:
            self.h2("Restrições e inconformidades registradas pelo órgão")
            self.legenda_tabela("Situação das restrições/inconformidades registradas anteriormente")
            self.tabela(["Código", "Descrição", "Situação na vistoria"], [(r.get("codigo", ""), r.get("descricao", ""), r.get("situacao", "")) for r in restr], larguras=[2.2, 10.0, 3.8], tam=8)
        if not anom:
            self.p("Não foram constatadas anomalias ou não conformidades nos sistemas vistoriados.")
            return
        self.h2("Quadro-resumo e planilha de grau de risco")
        self.legenda_tabela("Anomalias e não conformidades, com prioridade e grau de risco")
        linhas = [
            (a.get("id", ""), a.get("local", ""), a.get("titulo") or a.get("descricao", "")[:90], ORIGEM.get(a.get("origem"), a.get("origem", "")), str(a.get("prioridade", "") or ""), GRAU_RISCO.get(a.get("_grau") or "", ""))
            for a in anom
        ]
        self.tabela(["ID", "Local", "Anomalia", "Origem", "Prior.", "Grau"], linhas, larguras=[1.2, 3.0, 6.4, 2.9, 1.1, 1.6], tam=8, alinhamentos=["c", "l", "l", "l", "c", "c"])
        self.h2("Descrição detalhada")
        if self.d.get("agrupar_por_sistema", True):
            ordem = []
            for s in sits:
                if s.get("sistema") and s["sistema"] not in ordem:
                    ordem.append(s["sistema"])
            for a in anom:
                if (a.get("sistema") or "Outros") not in ordem:
                    ordem.append(a.get("sistema") or "Outros")
            sit_por = {s.get("sistema"): s for s in sits}
            for sistema in ordem:
                grupo = [a for a in anom if (a.get("sistema") or "Outros") == sistema]
                if not grupo:
                    continue
                self.sub3 = getattr(self, "sub3", 0) + 1
                self.doc.add_heading(f"{self.sec}.{self.sub}.{self.sub3}. {sistema}", level=3)
                s = sit_por.get(sistema)
                if s and s.get("descricao"):
                    self.p(f"**Situação:** {SITUACAO_SISTEMA.get(s.get('situacao'), s.get('situacao',''))}. {s['descricao']}")
                for a in grupo:
                    self._anomalia(a, nivel=4)
        else:
            for a in anom:
                self._anomalia(a)

    def analise(self):
        self.h1("Análise técnica")
        if self.d.get("analise"):
            self.texto(self.d["analise"])
        anom = self.d.get("anomalias") or []
        if self.tipo == "patologias" and anom:
            self.h2("Quadro de nexo causal")
            self.legenda_tabela("Síntese de causas, origens e nexo causal")
            self.tabela(
                ["ID", "Manifestação", "Causa", "Origem", "Nexo causal"],
                [(a.get("id", ""), a.get("titulo", ""), a.get("causa_resumo") or (a.get("causa", "")[:160]), ORIGEM.get(a.get("origem"), ""), a.get("nexo_resumo") or (a.get("nexo_causal", "")[:160])) for a in anom],
                larguras=[1.2, 3.6, 4.6, 2.6, 4.0], tam=8,
            )
        if self.tipo == "inspecao_predial":
            self.h2("Avaliação da manutenção")
            self.texto(self.d.get("avaliacao_manutencao") or "Não informada.")
            self.h2("Avaliação do uso")
            self.texto(self.d.get("avaliacao_uso") or "Não informada.")
        g = self.d.get("grau_risco_global")
        if g or anom:
            self.h2("Classificação global")
            cont = {p: sum(1 for a in anom if a.get("prioridade") == p) for p in (1, 2, 3)}
            txt = (
                f"Foram registradas {len(anom)} anomalias/não conformidades: {cont[1]} de prioridade 1 (crítico), "
                f"{cont[2]} de prioridade 2 (regular) e {cont[3]} de prioridade 3 (mínimo)."
            )
            if g:
                txt += f" A edificação é classificada, de forma global, com **grau de risco {GRAU_RISCO.get(g, g).lower()}**."
            if self.d.get("justificativa_grau_global"):
                txt += " " + self.d["justificativa_grau_global"]
            self.p(txt)
        if self.tipo == "recebimento_obra" and self.d.get("parecer_recebimento"):
            self.h2("Parecer quanto ao recebimento")
            self.texto(self.d["parecer_recebimento"])

    def custos(self):
        c = self.c
        tem_orc = bool(c["linhas"])
        eb = c["etapas_b"]
        if not tem_orc and not eb and not c["materiais"]:
            return
        self.h1("Quantificação e custos")
        base = self.d.get("base_custos") or {}
        if base:
            regime = "desonerada" if base.get("regime") == "desonerado" else "não desonerada"
            self.p(
                f"Os custos foram estimados com base na tabela **{base.get('fonte','')}** — {base.get('uf','')}, mês de referência "
                f"**{base.get('mes_referencia','')}**, versão {regime}, com BDI de **{pct(c['bdi'])}**. "
                "Os quantitativos foram levantados na vistoria e/ou extraídos da documentação, conforme indicado na planilha anexa."
            )
            for o in base.get("observacoes") or []:
                self.bullet(o)
        if tem_orc:
            t = c["totais"]
            self.h2("Orçamento discriminado")
            self.p(
                "Para cada serviço foram apuradas a quantidade prevista, a executada e aceita, a executada com vício ou danificada (a refazer) e, por diferença, a não executada. "
                "Serviço executado com vício ou deteriorado não é considerado executado; seu custo de refazimento inclui a remoção do serviço defeituoso e a reexecução."
            )
            if any(l["pu"] is None for l in c["linhas"]):
                self.p(f"**Atenção:** {t['sem_preco']} serviço(s) sem preço unitário definido; estão fora dos totais até o preenchimento na planilha.", italic=True)
            self.legenda_tabela("Resumo por etapa")
            linhas = [(e["etapa"], moeda(e["v_prev"]), moeda(e["v_aceito"]), pct(e["pct_aceito"]), moeda(e["c_refazer"]), moeda(e["c_executar"])) for e in c["por_etapa"]]
            total = ("TOTAL", moeda(t["v_prev"]), moeda(t["v_aceito"]), pct(t["pct_aceito"]), moeda(t["c_refazer"]), moeda(t["c_executar"]))
            self.tabela(["Etapa", "Valor previsto", "Executado aceito", "% aceito", "A refazer", "A executar"], linhas, larguras=[4.4, 2.5, 2.5, 1.6, 2.5, 2.5], tam=8, total=total, alinhamentos=["l", "r", "r", "r", "r", "r"])
            self.legenda_tabela("Quantitativos por serviço")
            ls = [
                (l["item"], f"{l['descricao']}{(' (' + l['fonte'] + ' ' + l['codigo'] + ')') if l['codigo'] else ''}", l["unidade"], num(l["prev"]), num(l["aceita"]), num(l["refazer"]), num(l["nao_exec"]), moeda(l["pu_bdi"]) if l["pu"] is not None else "a definir", moeda(l["c_concluir"]) if l["pu"] is not None else "a definir")
                for l in c["linhas"]
            ]
            self.tabela(["Item", "Serviço", "Un", "Prev.", "Aceita", "Refazer", "Não exec.", "PU c/ BDI", "Custo p/ concluir"], ls, larguras=[1.0, 4.8, 0.9, 1.3, 1.3, 1.3, 1.3, 1.9, 2.2], tam=7, alinhamentos=["c", "l", "c", "r", "r", "r", "r", "r", "r"])
            self.legenda_tabela("Síntese dos resultados")
            self.tabela(["Indicador", "Valor"], [
                ("Valor total previsto", moeda(t["v_prev"])),
                ("Valor executado e aceito", moeda(t["v_aceito"])),
                ("Percentual executado aceito", pct(t["pct_aceito"])),
                ("Percentual não executado ou a refazer", pct(1 - t["pct_aceito"]) if t["pct_aceito"] is not None else "—"),
                ("Percentual executado bruto (inclui serviços com vício)", pct(t["pct_bruto"])),
                ("Custo de refazer (serviços com vício ou danificados)", moeda(t["c_refazer"])),
                ("Custo a executar (não executado / pendências)", moeda(t["c_executar"])),
            ], larguras=[10.0, 6.0], tam=9, total=("Custo total para concluir/corrigir", moeda(t["c_concluir"])), alinhamentos=["l", "r"])
        if eb:
            self.h2("Percentual executado por etapas")
            cub = eb["cub"] or {}
            if eb["base"] == "contrato":
                base_txt = f"valor de contrato de {moeda(eb['contrato'])}"
            else:
                base_txt = f"CUB/m² {cub.get('projeto_padrao','')} ({cub.get('fonte','')}, {cub.get('mes','')}) de {moeda(float(cub['valor'])) if cub.get('valor') not in (None, '') else 'a definir'} × área equivalente de {num(eb['area'])} m²"
            if eb["adicional"]:
                base_txt += f", acrescido de {moeda(eb['adicional'])} referentes a itens não cobertos pelo CUB"
            self.p(f"Valor global de referência: {base_txt}, totalizando **{moeda(eb['valor_global'])}**. Pesos das etapas: {eb['pesos_origem'] or 'origem não informada'}.")
            self.legenda_tabela("Percentual executado por etapa")
            linhas = [(i["etapa"], pct(i["peso"]), pct(i["pct"]), pct(i["contrib"]), moeda(i["v_exec"])) for i in eb["itens"]]
            self.tabela(["Etapa", "Peso", "% executado", "Contribuição", "Valor executado"], linhas, larguras=[6.4, 2.0, 2.4, 2.4, 2.8], tam=8,
                        total=("TOTAL", pct(eb["soma_peso"]), "", pct(eb["pct_global"]), moeda(eb["v_exec"])), alinhamentos=["l", "r", "r", "r", "r"])
            self.p(f"Percentual executado pelo método das etapas: **{pct(eb['pct_global'])}**; valor a executar: **{moeda(eb['v_rest'])}**.")
        prog = self.d.get("programa") or {}
        if prog.get("percentual_informado") not in (None, ""):
            verificado = c["totais"]["pct_aceito"] if tem_orc and c["totais"]["pct_aceito"] is not None else (eb["pct_global"] if eb else None)
            inf = _frac(prog["percentual_informado"])
            self.h2("Confronto com o percentual informado")
            if verificado is not None:
                dif = verificado - inf
                self.p(
                    f"Percentual físico informado ({prog.get('fonte_percentual','fonte não indicada')}): **{pct(inf)}**. "
                    f"Percentual executado e aceito apurado neste laudo: **{pct(verificado)}** (diferença de {pct(dif)})."
                )
            if prog.get("explicacao_diferenca"):
                self.texto(prog["explicacao_diferenca"])
        if c["materiais"]:
            self.h2("Materiais entregues e não incorporados")
            self.legenda_tabela("Materiais no canteiro não incorporados à obra")
            self.tabela(["Material", "Un", "Qtd.", "Estado", "Valor"], [(m.get("descricao", ""), m.get("unidade", ""), num(m["quantidade"]), m.get("estado", ""), moeda(m["total"])) for m in c["materiais"]], larguras=[6.0, 1.2, 1.6, 4.4, 2.8], tam=8, alinhamentos=["l", "c", "r", "l", "r"])

    def recomendacoes(self):
        anom = self.d.get("anomalias") or []
        gerais = self.d.get("recomendacoes_gerais") or []
        if not anom and not gerais:
            return
        self.h1("Recomendações técnicas")
        if anom:
            ordem = sorted(anom, key=lambda a: (a.get("prioridade") or 9, str(a.get("id"))))
            self.legenda_tabela("Recomendações em ordem de prioridade")
            self.tabela(
                ["Prior.", "ID", "Recomendação", "Prazo sugerido"],
                [(str(a.get("prioridade", "") or ""), a.get("id", ""), a.get("recomendacao_resumo") or (a.get("recomendacao", "")[:220]), a.get("prazo", "")) for a in ordem],
                larguras=[1.3, 1.3, 10.4, 3.0], tam=8, alinhamentos=["c", "c", "l", "l"],
            )
        for g in gerais:
            self.bullet(g)
        if self.tipo == "inspecao_predial" and self.d.get("periodicidade_proxima_inspecao"):
            self.p(f"**Próxima inspeção:** {self.d['periodicidade_proxima_inspecao']}")

    def quesitos(self):
        qs = self.d.get("quesitos") or []
        if not qs:
            return
        self.h1("Respostas aos quesitos")
        parte_atual = None
        for q in qs:
            if q.get("parte") and q["parte"] != parte_atual:
                parte_atual = q["parte"]
                self.h2(f"Quesitos — {parte_atual}")
            self.p(f"**Quesito {q.get('numero','')}:** {q.get('pergunta','')}", after=2)
            self.p(f"**Resposta:** {q.get('resposta','')}")

    def conclusao(self):
        self.h1("Conclusão")
        if self.d.get("conclusao"):
            self.texto(self.d["conclusao"])
        else:
            self.avisos.append("Campo 'conclusao' vazio — o laudo precisa de conclusão redigida.")
            self.p("[Conclusão a redigir]", italic=True)

    def encerramento(self):
        self.h1("Encerramento")
        n_fotos = self._contar_fotos()
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        par.add_run("Nada mais havendo a relatar, encerra-se o presente laudo, composto por ")
        r = par.add_run()
        _campo(r, "NUMPAGES")
        par.add_run(f" páginas numeradas{f' e {n_fotos} fotografias' if n_fotos else ''}, acompanhado da planilha de quantitativos e custos e dos anexos relacionados.")
        anexos = ["ART do laudo"] + list(self.d.get("anexos") or [])
        if self.c["linhas"] or self.c["etapas_b"]:
            anexos.insert(1, f"Planilha de quantitativos e custos (Planilha_{self.d.get('numero_laudo','')}.xlsx)")
        if self.d.get("fotos_gerais"):
            anexos.append("Relatório fotográfico")
        self.p("**Anexos:**", after=2)
        for i, a in enumerate(anexos, 1):
            self.bullet(f"Anexo {i} — {a}")
        self.doc.add_paragraph()
        self.p(f"{self.d.get('local_emissao','')}, {data_extenso(self.d.get('data_emissao',''))}.", align="center")
        for _ in range(3):
            self.doc.add_paragraph()
        signatarios = [self.perito] + list(self.d.get("coautores") or [])
        for s in signatarios:
            self.p("_______________________________________________", align="center", after=0)
            self.p(s.get("nome", ""), bold=True, align="center", after=0)
            self.p(s.get("titulo", ""), align="center", after=0)
            regs = " | ".join([s.get("registro", "")] + list(s.get("registros_adicionais") or []))
            self.p(regs, align="center", after=0)
            if s is self.perito:
                self.p(f"ART nº {self.perito.get('art') or '________________'}", align="center", after=0)
            self.doc.add_paragraph()

    def _contar_fotos(self):
        n = 0
        for a in self.d.get("anomalias") or []:
            n += len([f for f in a.get("fotos") or [] if f.get("arquivo")])
        n += len([f for f in self.d.get("fotos_gerais") or [] if f.get("arquivo")])
        return n

    def anexo_fotos(self):
        fg = self.d.get("fotos_gerais") or []
        if not fg:
            return
        self.doc.add_page_break()
        self.doc.add_heading("ANEXO — RELATÓRIO FOTOGRÁFICO", level=1)
        grupos = []
        for f in fg:
            g = f.get("grupo") or ""
            if g not in grupos:
                grupos.append(g)
        for g in grupos:
            if g:
                self.doc.add_heading(g, level=2)
            self.fotos([f for f in fg if (f.get("grupo") or "") == g])

    # ---------------------------------------------------------- montagem
    def gerar(self, caminho):
        self.capa()
        self.sumario()
        self.identificacao()
        self.objeto()
        self.dados_obra()
        self.pressupostos()
        self.normas()
        self.metodologia()
        self.documentos()
        self.caracterizacao()
        self.vistoria()
        self.constatacoes()
        self.analise()
        self.custos()
        self.recomendacoes()
        self.quesitos()
        self.conclusao()
        self.encerramento()
        self.anexo_fotos()
        if self._toc_run is not None:
            for i, (nivel, titulo) in enumerate(self.toc):
                if i:
                    self._toc_run.add_break()
                self._toc_run.add_text(("      " if nivel == 2 else "") + titulo)
            self._toc_run.font.size = Pt(10)
        self.doc.save(caminho)
        return self.avisos


def _frac(v):
    try:
        v = float(str(v).replace("%", "").replace(",", "."))
    except ValueError:
        return None
    return v / 100 if v > 1 else v
